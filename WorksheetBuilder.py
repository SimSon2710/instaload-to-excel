import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Alignment
from datetime import datetime
from InstaPost import InstaPost


class Worksheet:
    WB = openpyxl.Workbook()

    WSPosts = WB.worksheets[0]
    WSPosts.title = "Posts"
    if InstaPost.comments_count > 0:
        WB.create_sheet("PostComments", 1)
        WSComments = WB.worksheets[1]
    else:
        WSComments = None

    col_titles_posts = []
    col_titles_comments = []

    col_specs_posts = {
        'A': {'width': 20, 'title': 'profile_name', 'wrap_text': False, 'var_name': 'username'},
        'B': {'width': 50, 'title': 'post_text', 'wrap_text': True, 'var_name': 'text'},
        'C': {'width': 30, 'title': 'post_image', 'wrap_text': False, 'var_name': ''},
        'D': {'width': 23, 'title': 'post_time', 'wrap_text': False, 'var_name': 'post_time'},
        'E': {'width': 15, 'title': 'post_likes', 'wrap_text': False, 'var_name': 'post_likes'},
        'F': {'width': 20, 'title': 'post_id', 'wrap_text': False, 'var_name': 'post_id'},
        'G': {'width': 20, 'title': 'profile_full_name', 'wrap_text': False,
              'var_name': 'profile_full_name'},
        'H': {'width': 20, 'title': 'profile_category', 'wrap_text': False,
              'var_name': 'profile_category'},
        'I': {'width': 30, 'title': 'profile_biography', 'wrap_text': True,
              'var_name': 'profile_biography'},
        'J': {'width': 14, 'title': 'profile_followed_by', 'wrap_text': False,
              'var_name': 'profile_followed_by'},
        'K': {'width': 12, 'title': 'profile_follows', 'wrap_text': False,
              'var_name': 'profile_follows'},
        'L': {'width': 23, 'title': 'profile_is_business_account', 'wrap_text': False,
              'var_name': 'profile_is_business_account'},
        'M': {'width': 21, 'title': 'profile_is_joined_recently', 'wrap_text': False,
              'var_name': 'profile_is_joined_recently'},
        'N': {'width': 30, 'title': 'profile_external_url', 'wrap_text': False,
              'var_name': 'profile_external_url'},
        'O': {'width': 23, 'title': 'post_image_is_part_of_collection', 'wrap_text': False,
              'var_name': 'post_is_part_of_collection'}
    }

    col_specs_comments = {
        'A': {'width': 20, 'title': 'post_id', 'wrap_text': False, 'var_name': 'post_id'},
        'B': {'width': 20, 'title': 'post_of_user', 'wrap_text': False, 'var_name': 'post_of_user'},
        'C': {'width': 20, 'title': 'comment_id', 'wrap_text': False, 'var_name': 'id'},
        'D': {'width': 30, 'title': 'creation_time', 'wrap_text': False, 'var_name': 'created_at'},
        'E': {'width': 50, 'title': 'comment_text', 'wrap_text': True, 'var_name': 'text'},
        'F': {'width': 15, 'title': 'user_name', 'wrap_text': False, 'var_name': 'username'},
        'G': {'width': 20, 'title': 'user_id', 'wrap_text': False, 'var_name': 'id'},
        'H': {'width': 20, 'title': 'is_answer', 'wrap_text': False,
              'var_name': 'answers'},
        'I': {'width': 20, 'title': 'parent_comment_id', 'wrap_text': False,
              'var_name': 'id'}
    }

    for col in col_specs_posts:
        col_titles_posts.append(col_specs_posts[col]['title'])
        WSPosts.column_dimensions[col].width = col_specs_posts[col]['width']
    WSPosts.append(col_titles_posts)

    if WSComments:
        for col in col_specs_comments:
            col_titles_comments.append(col_specs_comments[col]['title'])
            WSComments.column_dimensions[col].width = col_specs_comments[col]['width']
        WSComments.append(col_titles_comments)

    cur_row_posts = 2
    cur_row_comments = 2

    def __init__(self, posts: dict):
        self.Posts = posts
        self.row_height = 185

    def add_to_post_ws(self):
        skip = ['post_image', 'post_image_file_name']
        for Image in self.Posts:
            post = self.Posts[Image]
            for col in self.col_specs_posts:
                if not self.col_specs_posts[col]['title'] in skip:
                    post_content = getattr(post, self.col_specs_posts[col]['var_name'])
                    self.WSPosts[f'{col}{str(self.cur_row_posts)}'] = post_content
                    self.WSPosts[f'{col}{str(self.cur_row_posts)}'].alignment = \
                        Alignment(wrap_text=self.col_specs_posts[col]['wrap_text'])
                elif self.col_specs_posts[col]['title'] == 'post_image':
                    Image.anchor = f'{col}{str(self.cur_row_posts)}'
                    self.WSPosts.add_image(Image)
            self.WSPosts.row_dimensions[self.cur_row_posts].height = post.row_height
            self.cur_row_posts += 1

    def add_to_comments_ws(self):
        if Worksheet.WSComments:
            cur_post_id = None
            for Image in self.Posts:
                Post = self.Posts[Image]
                if Post.comments and not Post.comments["post_id"] == cur_post_id:
                    cur_post_id = Post.comments["post_id"]
                    comments = Post.comments
                    for comment in comments["comments"]:
                        for col in self.col_specs_comments:
                            has_answers = True if comment['answers'] else False
                            content = None
                            if self.col_specs_comments[col]["title"] in ["comment_id", "comment_text",
                                                                         "user_id"]:
                                content = comment[self.col_specs_comments[col]["var_name"]]
                            if self.col_specs_comments[col]["title"] == "creation_time":
                                content = datetime.utcfromtimestamp(int(comment[self.col_specs_comments[col]["var_name"]]))\
                                    .strftime("%Y-%m-%d %H:%M:%S")
                            if self.col_specs_comments[col]["title"] == "user_name":
                                content = comment["owner"][self.col_specs_comments[col]["var_name"]]
                            if self.col_specs_comments[col]["title"] in ["post_id", "post_of_user"]:
                                content = comments[self.col_specs_comments[col]["var_name"]]
                            if has_answers:
                                if self.col_specs_comments[col]["title"] == "parent_comment_id":
                                    content = comment["id"]
                                elif self.col_specs_comments[col]["title"] == "is_answer":
                                    content = "True"
                            else:
                                if self.col_specs_comments[col]["title"] == "is_answer":
                                    content = "False"

                            self.WSComments[f'{col}{str(self.cur_row_comments)}'] = str(content)
                            self.WSComments[f'{col}{str(self.cur_row_comments)}'].alignment = \
                                Alignment(wrap_text=self.col_specs_comments[col]['wrap_text'])

                        self.cur_row_comments += 1

    def create_workbook(self, save_path):
        tab = Table(displayName="Posts", ref=f"A1:{list(self.col_specs_posts)[-1]}{str(self.cur_row_posts)}")
        tab_style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                                   showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        tab.tableStyleInfo = tab_style
        self.WSPosts.add_table(tab)

        if Worksheet.WSComments:
            tab_comments = Table(displayName="Comments",
                                 ref=f"A1:{list(self.col_specs_comments)[-1]}{str(self.cur_row_comments)}")
            tab_style_comments = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                                                showLastColumn=False, showRowStripes=True, showColumnStripes=True)
            tab_comments.tableStyleInfo = tab_style_comments
            self.WSComments.add_table(tab_comments)

        self.WB.save(save_path)
